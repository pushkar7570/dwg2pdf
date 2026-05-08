
#!/usr/bin/env python3
"""
control_design_assessment.py
---------------------------------
Create a lightweight, repeatable **design assessment** testing package for a control.
Covers Who, What, When, Where, Why, How, and How Often (frequency), and produces:
 - A Markdown test plan (design procedures + acceptance criteria)
 - An evidence/artifacts checklist (CSV and, if openpyxl is installed, XLSX)
 - A normalized JSON profile of the control

Usage examples:
  # 1) Start interactively
  python control_design_assessment.py --interactive

  # 2) From a JSON/YAML profile
  python control_design_assessment.py --input control_profile.json --outdir out/

  # 3) Generate a blank starter template you can fill in
  python control_design_assessment.py --init-template my_control_template.json

Notes:
- YAML input is supported if PyYAML is installed (pip install pyyaml).
- XLSX export is supported if openpyxl is installed (pip install openpyxl).
- This script focuses on **design** assessment (suitability of design and placement in operation),
  not operating effectiveness testing (sampling over a period).
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# ---------- Data Model ----------

@dataclass
class ControlProfile:
    # Core identifiers
    control_id: str
    title: str
    process_area: str
    risk_statement: str
    objective: str

    # 5W + H + Frequency
    who: Dict[str, List[str]]  # {"owner": [...], "performer": [...], "reviewer": [...]}
    what: str
    when: str  # timing/trigger
    where: str  # systems/locations
    why: str  # link to objective/risk
    how: str  # method/steps/criteria
    frequency: str  # e.g., "daily", "monthly", "per transaction", "continuous", etc.

    # Additional design attributes
    control_type: str  # preventive/detective/corrective + manual/automated/MRE
    ipe: List[Dict[str, str]] = field(default_factory=list)  # [{'name','source','fields','owner'}]
    precision: str = ""  # thresholds, review criteria, exception handling
    evidence_retention: str = ""  # retention policy/period
    segregation_of_duties: str = ""  # SoD considerations
    population_definition: str = ""  # what population would exist if operated
    artifacts_needed: List[str] = field(default_factory=list)

    # Metadata
    last_updated: str = field(default_factory=lambda: datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"))
    notes: str = ""


# ---------- Helpers ----------

FREQUENCY_NORMALIZER = {
    "ad hoc": "ad hoc",
    "adhoc": "ad hoc",
    "as-needed": "ad hoc",
    "continuous": "continuous",
    "real-time": "continuous",
    "per transaction": "per transaction",
    "transactional": "per transaction",
    "daily": "daily",
    "weekly": "weekly",
    "biweekly": "biweekly",
    "monthly": "monthly",
    "quarterly": "quarterly",
    "semi-annual": "semi-annual",
    "semiannual": "semi-annual",
    "annual": "annual",
}

def normalize_frequency(freq: str) -> str:
    f = (freq or "").strip().lower()
    # try exact match
    if f in FREQUENCY_NORMALIZER:
        return FREQUENCY_NORMALIZER[f]
    # try loose matching
    for k in FREQUENCY_NORMALIZER:
        if k in f:
            return FREQUENCY_NORMALIZER[k]
    return f or "unspecified"

def safe_list(val: Any) -> List[str]:
    if val is None:
        return []
    if isinstance(val, list):
        return [str(x) for x in val if str(x).strip()]
    return [str(val)]

def ensure_roles(who: Dict[str, List[str]]) -> Dict[str, List[str]]:
    who = who or {}
    return {
        "owner": safe_list(who.get("owner")),
        "performer": safe_list(who.get("performer")),
        "reviewer": safe_list(who.get("reviewer")),
    }

def suggest_artifacts(profile: ControlProfile) -> List[str]:
    """Very light heuristics to propose artifacts if not provided."""
    suggestions = set()
    text_blob = " ".join([profile.what, profile.where, profile.how]).lower()

    # Universal design artifacts
    suggestions.update([
        "Policy/Procedure document",
        "Control narrative / flowchart",
        "Walkthrough meeting notes",
    ])

    keywords = {
        "access": ["access listing/export", "provisioning ticket(s)", "termination ticket(s)", "periodic access review report", "role/permission matrix", "SoD ruleset/report"],
        "change": ["change ticket(s)", "pull request / code review evidence", "deployment logs", "segregation between DEV/QA/PROD proof"],
        "configuration": ["configuration export / screenshots", "parameter setting report", "baseline configuration"],
        "logging": ["system audit logs", "SIEM alerts", "retention/archival settings"],
        "reconcile": ["reconciliation report", "variance/exceptions log", "evidence of investigation & resolution"],
        "approval": ["approval email or workflow record", "threshold/criteria documentation"],
        "ticket": ["ticketing system report (Jira/ServiceNow)", "assignment & closure evidence"],
        "job": ["batch job schedule", "job run logs", "failure alerts & resolutions"],
        "report": ["source report used by control (IPE)", "report definition / SQL", "data lineage / field mapping"],
        "security": ["user/role listing", "MFA settings", "password policy configuration"],
    }

    for key, arts in keywords.items():
        if key in text_blob:
            suggestions.update(arts)

    # If IPE is listed, ensure completeness/accuracy artifacts
    if profile.ipe:
        suggestions.update([
            "IPE definition/specification",
            "Report parameters & filter screenshots",
            "Evidence of IPE completeness & accuracy (CA) testing",
        ])

    # If manual review exists
    if re.search(r"\b(review|approve|analy[sz]e|investigate)\b", text_blob):
        suggestions.update([
            "Reviewer sign-off (timestamped)",
            "Evidence of review criteria/threshold applied",
            "Exception tracker & follow-up evidence",
        ])

    # If automated
    if "automated" in (profile.control_type or "").lower():
        suggestions.update(["System configuration proof", "Change management evidence for automation"])

    # Frequency-based
    freq = normalize_frequency(profile.frequency)
    if freq in {"daily", "weekly", "monthly", "quarterly", "annual", "semi-annual"}:
        suggestions.update(["Schedule/calendar entry", "Evidence of timely performance"])
    if freq in {"per transaction", "continuous"}:
        suggestions.update(["Event trigger proof", "Exception alert sample"])

    # Always include evidence management
    suggestions.update(["Evidence retention policy reference"])

    # Combine with any provided
    final = list(dict.fromkeys([*profile.artifacts_needed, *sorted(suggestions)]))
    return final


# ---------- Loaders ----------

def load_profile(path: Path) -> Dict[str, Any]:
    ext = path.suffix.lower()
    text = path.read_text(encoding="utf-8")
    if ext in {".json"}:
        return json.loads(text)
    if ext in {".yaml", ".yml"}:
        try:
            import yaml  # type: ignore
        except Exception as e:
            raise RuntimeError("YAML input requested but PyYAML is not installed. pip install pyyaml") from e
        return yaml.safe_load(text)
    if ext in {".csv"}:
        # Expect two-column key,value CSV or a header-based single-row CSV
        rows = list(csv.DictReader(text.splitlines()))
        if len(rows) == 1:
            return rows[0]
        d: Dict[str, Any] = {}
        for r in rows:
            key = r.get("key") or r.get("Key") or r.get("field")
            val = r.get("value") or r.get("Value")
            if key:
                d[str(key)] = val
        return d
    raise ValueError(f"Unsupported input format: {ext}")


def prompt_multiline(prompt: str) -> str:
    print(f"{prompt} (finish with a single '.' on its own line):")
    lines = []
    while True:
        line = input()
        if line.strip() == ".":
            break
        lines.append(line)
    return "\n".join(lines).strip()


def interactive_collect() -> Dict[str, Any]:
    print("\n=== Interactive Control Design Profile ===\n")
    d: Dict[str, Any] = {}
    d["control_id"]      = input("Control ID: ").strip()
    d["title"]           = input("Title: ").strip()
    d["process_area"]    = input("Process Area (e.g., Access Mgmt, Change Mgmt, Financial Close): ").strip()
    d["risk_statement"]  = prompt_multiline("Risk Statement")
    d["objective"]       = prompt_multiline("Control Objective")

    print("\n-- Who performs / owns the control --")
    d["who"] = {
        "owner": [s.strip() for s in input("Owner(s) (comma separated names/roles): ").split(",") if s.strip()],
        "performer": [s.strip() for s in input("Performer(s) (comma separated names/roles): ").split(",") if s.strip()],
        "reviewer": [s.strip() for s in input("Reviewer(s) (comma separated names/roles): ").split(",") if s.strip()],
    }

    d["what"]            = prompt_multiline("WHAT (describe the activity / nature of the control)")
    d["when"]            = input("WHEN (timing/trigger): ").strip()
    d["where"]           = input("WHERE (systems/locations): ").strip()
    d["why"]             = prompt_multiline("WHY (link to objective/risk)")
    d["how"]             = prompt_multiline("HOW (method/steps/criteria applied; include thresholds)")
    d["frequency"]       = input("HOW OFTEN (frequency): ").strip()
    d["control_type"]    = input("Control type (e.g., preventive/detective, manual/automated/MRE): ").strip()

    print("\n-- IPE (Information Produced by the Entity) used by the control --")
    ipe_items: List[Dict[str, str]] = []
    while True:
        add = input("Add an IPE item? (y/n): ").strip().lower()
        if add != "y":
            break
        name = input("  Name of report/log/data object: ").strip()
        source = input("  Source system: ").strip()
        fields = input("  Key fields/parameters: ").strip()
        owner  = input("  Data/report owner: ").strip()
        ipe_items.append({"name": name, "source": source, "fields": fields, "owner": owner})
    d["ipe"] = ipe_items

    d["precision"]             = prompt_multiline("Precision (review criteria/thresholds, what constitutes an exception)")
    d["evidence_retention"]    = input("Evidence retention (policy/period): ").strip()
    d["segregation_of_duties"] = prompt_multiline("Segregation of duties considerations (if applicable)")
    d["population_definition"] = prompt_multiline("Population definition (what would constitute a population for OE)")

    arts = input("List any known artifacts/evidence to collect (comma separated) or leave blank: ").strip()
    d["artifacts_needed"] = [s.strip() for s in arts.split(",") if s.strip()] if arts else []

    d["notes"] = prompt_multiline("Additional notes")
    return d


# ---------- Generators ----------

def make_profile(d: Dict[str, Any]) -> ControlProfile:
    who = ensure_roles(d.get("who", {}))
    profile = ControlProfile(
        control_id = str(d.get("control_id", "")).strip() or "TBD",
        title = str(d.get("title", "")).strip() or "TBD Title",
        process_area = str(d.get("process_area", "")).strip(),
        risk_statement = str(d.get("risk_statement", "")).strip(),
        objective = str(d.get("objective", "")).strip(),
        who = who,
        what = str(d.get("what", "")).strip(),
        when = str(d.get("when", "")).strip(),
        where = str(d.get("where", "")).strip(),
        why = str(d.get("why", "")).strip(),
        how = str(d.get("how", "")).strip(),
        frequency = normalize_frequency(str(d.get("frequency", "")).strip()),
        control_type = str(d.get("control_type", "")).strip(),
        ipe = [{
            "name": str(x.get("name","")).strip(),
            "source": str(x.get("source","")).strip(),
            "fields": str(x.get("fields","")).strip(),
            "owner": str(x.get("owner","")).strip(),
        } for x in d.get("ipe", []) if isinstance(x, dict)],
        precision = str(d.get("precision","")).strip(),
        evidence_retention = str(d.get("evidence_retention","")).strip(),
        segregation_of_duties = str(d.get("segregation_of_duties","")).strip(),
        population_definition = str(d.get("population_definition","")).strip(),
        artifacts_needed = [str(a).strip() for a in d.get("artifacts_needed", []) if str(a).strip()],
        notes = str(d.get("notes","")).strip(),
    )
    # Expand artifacts
    profile.artifacts_needed = suggest_artifacts(profile)
    return profile


def render_summary_table(profile: ControlProfile) -> str:
    def list_or_dash(items: List[str]) -> str:
        return ", ".join(items) if items else "—"
    return f"""
| Attribute | Details |
|---|---|
| **Who** | Owner: {list_or_dash(profile.who.get('owner'))}; Performer: {list_or_dash(profile.who.get('performer'))}; Reviewer: {list_or_dash(profile.who.get('reviewer'))} |
| **What** | {profile.what or '—'} |
| **When** | {profile.when or '—'} |
| **Where** | {profile.where or '—'} |
| **Why** | {profile.why or '—'} |
| **How** | {profile.how or '—'} |
| **How Often (Frequency)** | {profile.frequency or '—'} |
| **Control Type** | {profile.control_type or '—'} |
| **Evidence Retention** | {profile.evidence_retention or '—'} |
| **Segregation of Duties** | {profile.segregation_of_duties or '—'} |
| **Population Definition (for context)** | {profile.population_definition or '—'} |
""".strip()


def render_procedures(profile: ControlProfile) -> str:
    """Design assessment procedures (non-statistical, suitability of design)."""
    steps = [
        ("Understand the control and risk coverage",
         "Read the policy/procedure, control narrative/flowchart, and risk & control matrix entries. "
         "Confirm that the stated **Why** aligns with the risk statement and objective, and that the **How** and **What** would address the risk if performed as designed.",
         "Control activity addresses the stated risk/objective without critical gaps or conflicting responsibilities."),

        ("Walkthrough with control owner/performer",
         "Interview the **Who** (owner/performer/reviewer) to understand the **When**, **Where**, **How**, and **frequency**. "
         "Trace at least one recent instance (if placed in operation) end‑to‑end to observe the steps and information used (IPE).",
         "Roles are appropriate, steps are understood and consistently described, and the walkthrough demonstrates feasibility of execution."),

        ("Evaluate precision and criteria",
         "Inspect review thresholds/criteria in the **How** and confirm they are specific and capable of identifying material errors or noncompliance. "
         "For automated controls, confirm configurations/parameters. For MRE/manual reviews, confirm exception handling and follow‑up requirements.",
         "Defined criteria/thresholds are sufficiently precise; exceptions and follow‑ups are required and tracked."),

        ("Assess IPE (Information Produced by the Entity) reliance",
         "Identify each IPE item used (reports/logs/exports). For each, obtain definition, parameters, and source. "
         "Determine what procedures would be necessary to gain comfort over IPE completeness & accuracy (e.g., tie‑out key fields to source, parameters screenshots).",
         "IPE items are clearly defined, owned, and procedures exist (or can be performed) to support completeness & accuracy."),

        ("Check segregation of duties and access",
         "Evaluate whether performers/reviewers have conflicting access (e.g., ability to both make and approve changes). Review SoD considerations and evidence.",
         "No unmitigated SoD conflicts exist; if conflicts exist, mitigating controls are identified and designed appropriately."),

        ("Consider frequency and timeliness expectations",
         "Verify that **How Often** aligns with the risk exposure (e.g., daily vs. monthly). Confirm evidence of timeliness (timestamps) and schedules/calendars where applicable.",
         "Frequency is appropriate for the risk; timeliness expectations are defined and evidenced."),

        ("Confirm evidence sufficiency & retention",
         "Inspect whether evidence to prove performance exists and is retained for the required period (e.g., sign‑offs, logs, tickets, configurations).",
         "Evidence is objective, tamper‑resistant, and retained per policy; sufficient to re‑perform/review after the fact."),
    ]

    md = ["\n## Design Assessment Procedures\n"]
    for i, (title, action, criteria) in enumerate(steps, start=1):
        md.append(f"**Step {i}: {title}**\n\n- Procedures: {action}\n- Acceptance criteria: {criteria}\n")
    return "\n".join(md).strip()


def render_artifacts_section(profile: ControlProfile) -> str:
    arts = profile.artifacts_needed or []
    lines = ["\n## Artifacts & Evidence Checklist (Design)\n",
             "Below are suggested items to obtain/inspect during the design assessment."]
    for idx, a in enumerate(arts, start=1):
        lines.append(f"- [ ] {a}")
    return "\n".join(lines).strip()


def render_ipe_table(profile: ControlProfile) -> str:
    if not profile.ipe:
        return ""
    lines = ["\n## IPE Register (Information Produced by the Entity)\n",
             "| Name | Source System | Key Fields/Parameters | Owner |",
             "|---|---|---|---|"]
    for item in profile.ipe:
        lines.append(f"| {item.get('name','')} | {item.get('source','')} | {item.get('fields','')} | {item.get('owner','')} |")
    return "\n".join(lines).strip()


def generate_test_plan_md(profile: ControlProfile) -> str:
    header = f"# Control Design Assessment – {profile.control_id}: {profile.title}\n\n" \
             f"**Process Area:** {profile.process_area}\n\n" \
             f"**Prepared:** {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}\n"
    summary = render_summary_table(profile)
    procedures = render_procedures(profile)
    artifacts = render_artifacts_section(profile)
    ipe = render_ipe_table(profile)
    conclusion = dedent("""
    ## Conclusion
    - Design evaluation result: ☐ Sufficiently designed  ☐ Partially designed  ☐ Not sufficiently designed
    - Key gaps noted (if any): …
    - Remediation recommendations: …
    """).strip()

    risk_obj = "\n## Risk & Objective\n" + f"**Risk Statement:**\n{profile.risk_statement or '—'}\n\n**Control Objective:**\n{profile.objective or '—'}\n"

    notes = ""
    if profile.notes:
        notes = "\n## Notes\n" + profile.notes + "\n"

    sections = [header, risk_obj, "## Summary (5W+H+Frequency)\n" + summary, procedures, artifacts, ipe, conclusion, notes]
    return "\n\n".join([s for s in sections if s]).strip()


def write_markdown(md: str, outdir: Path, control_id: str) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    path = outdir / f"{sanitize_filename(control_id)}_design_test_plan.md"
    path.write_text(md, encoding="utf-8")
    return path


def sanitize_filename(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name.strip() or "control")


def write_json(profile: ControlProfile, outdir: Path) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    path = outdir / f"{sanitize_filename(profile.control_id)}_profile.json"
    path.write_text(json.dumps(asdict(profile), indent=2), encoding="utf-8")
    return path


def write_evidence_checklist(profile: ControlProfile, outdir: Path) -> Tuple[Path, Optional[Path]]:
    outdir.mkdir(parents=True, exist_ok=True)
    base = sanitize_filename(profile.control_id) or "control"
    csv_path = outdir / f"{base}_evidence_checklist.csv"
    headers = ["Item #", "Artifact / Evidence", "Source System", "Owner/Point of Contact", "Obtained (Y/N)", "File Path/Link", "Notes"]
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for idx, art in enumerate(profile.artifacts_needed or [], start=1):
            w.writerow([idx, art, "", "", "", "", ""])

    xlsx_path = None
    try:
        import openpyxl  # type: ignore
        from openpyxl.styles import Font, Alignment

        xlsx_path = outdir / f"{base}_evidence_checklist.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Checklist"
        ws.append(headers)
        bold = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, art in enumerate(profile.artifacts_needed or [], start=1):
            ws.append([idx, art, "", "", "", "", ""])
        # Column widths
        widths = [8, 50, 20, 28, 14, 40, 40]
        for i, wth in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wth
        wb.save(xlsx_path)
    except Exception:
        # openpyxl not available or failed; CSV already written
        xlsx_path = None

    return csv_path, xlsx_path


def generate_template(path: Path) -> None:
    template = {
        "control_id": "ITAC-001",
        "title": "Periodic Access Review for Key Financial Systems",
        "process_area": "Access Management",
        "risk_statement": "Unauthorized or excessive access may enable inappropriate transactions or data changes.",
        "objective": "Ensure user access remains appropriate and aligned with job responsibilities.",
        "who": {"owner": ["IT Security Manager"], "performer": ["System Owners"], "reviewer": ["Compliance Lead"]},
        "what": "Review active users and entitlements; identify exceptions; obtain remediation.",
        "when": "Monthly; triggered on the first business day following month-end",
        "where": "ERP, CRM, and Data Warehouse systems",
        "why": "To mitigate risk of unauthorized access to financial reporting systems.",
        "how": "Generate access listings from each system; compare to HR roster; reviewers attest to appropriateness; exceptions tracked to remediation. Threshold: all privileged roles reviewed 100%.",
        "frequency": "Monthly",
        "control_type": "Detective - Manual Review (MRE)",
        "ipe": [
            {"name": "ERP User Access Export", "source": "ERP", "fields": "user_id, role, last_login", "owner": "ERP Admin"},
            {"name": "HR Active Employee Roster", "source": "HRIS", "fields": "employee_id, status, department", "owner": "HRIS Analyst"}
        ],
        "precision": "All privileged roles reviewed 100%; any orphan accounts remediated within 5 business days.",
        "evidence_retention": "1 year",
        "segregation_of_duties": "Reviewers do not have provisioning rights; mitigating control: ticket workflow approval.",
        "population_definition": "All active users with access to in-scope systems as of period-end.",
        "artifacts_needed": [
            "Policy/Procedure document",
            "Control narrative / flowchart",
            "ERP user access export",
            "HR active employee roster",
            "Reviewer sign-offs",
            "Exception tracker & follow-up evidence",
            "Evidence retention policy reference"
        ],
        "notes": "This is a sample; customize as needed."
    }
    path.write_text(json.dumps(template, indent=2), encoding="utf-8")
    print(f"Template written to: {path}")


# ---------- CLI ----------

def main():
    ap = argparse.ArgumentParser(description="Generate a control design assessment testing package (5W+H+Frequency + artifacts).")
    ap.add_argument("--input", type=str, help="Path to control profile (JSON/YAML/CSV).")
    ap.add_argument("--interactive", action="store_true", help="Collect inputs interactively in the terminal.")
    ap.add_argument("--outdir", type=str, default="out", help="Output directory (default: ./out)")
    ap.add_argument("--init-template", type=str, help="Write a starter JSON template to this path and exit.")
    args = ap.parse_args()

    if args.init_template:
        generate_template(Path(args.init_template))
        return

    if not args.input and not args.interactive:
        ap.error("Provide --interactive or --input <file>; or use --init-template to create a starter file.")

    if args.input and args.interactive:
        ap.error("Please choose either --interactive or --input, not both.")

    raw: Dict[str, Any]
    if args.interactive:
        raw = interactive_collect()
    else:
        inpath = Path(args.input)
        if not inpath.exists():
            raise FileNotFoundError(f"Input file not found: {inpath}")
        raw = load_profile(inpath)

    profile = make_profile(raw)

    outdir = Path(args.outdir)
    md_path = write_markdown(generate_test_plan_md(profile), outdir, profile.control_id)
    csv_path, xlsx_path = write_evidence_checklist(profile, outdir)
    json_path = write_json(profile, outdir)

    print("\n=== Outputs ===")
    print(f"Test plan (Markdown): {md_path}")
    print(f"Evidence checklist (CSV): {csv_path}")
    if xlsx_path:
        print(f"Evidence checklist (XLSX): {xlsx_path} (requires 'openpyxl')")
    else:
        print("Evidence checklist (XLSX): not created (install 'openpyxl' to enable)")
    print(f"Normalized profile (JSON): {json_path}")
    print("\nDone.")


if __name__ == "__main__":
    main()
